from __future__ import annotations

import io
import asyncio
from uuid import UUID
from typing import List, Optional, Dict, Tuple, Any

from sqlalchemy import select, update, insert, func
from sqlalchemy.orm import selectinload
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.benchmark_run import BenchmarkRun, RunStatus
from app.models.benchmark_execution import BenchmarkExecution, ExecutionStatus
from app.models.llm_model import LLMModel
from app.models.test_case import TestCase
from app.models.test_suite import TestSuite
from app.schemas.benchmark import BenchmarkRunCreate

try:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Font, Border, Side
except ImportError:  # pragma: no cover
    Workbook = None  # type: ignore


class BenchmarkService:
    async def create_run(self, db: AsyncSession, payload: BenchmarkRunCreate) -> Tuple[BenchmarkRun, int]:
        stmt = select(TestCase.id).where(TestCase.suite_id.in_(payload.suite_ids))
        result = await db.execute(stmt)
        test_case_ids = result.scalars().all()

        if not test_case_ids:
            raise ValueError("Wybrane zbiory testowe są puste (brak przypadków testowych).")

        default_name = f"Benchmark: {len(payload.model_ids)} modeli, {len(payload.suite_ids)} zbiorów"
        new_run = BenchmarkRun(
            name=payload.name or default_name,
            status=RunStatus.PENDING
        )

        db.add(new_run)
        await db.flush()

        executions_data = [
            {
                "run_id": new_run.id,
                "test_case_id": tc_id,
                "llm_model_id": model_id,
                "status": ExecutionStatus.PENDING
            }
            for model_id in payload.model_ids
            for tc_id in test_case_ids
        ]

        await db.execute(insert(BenchmarkExecution).values(executions_data))
        await db.commit()

        return new_run, len(executions_data)

    async def get_all_runs(self, db: AsyncSession) -> List[Any]:
        stmt = (
            select(
                BenchmarkRun,
                func.count(BenchmarkExecution.id).label("total_executions")
            )
            .outerjoin(BenchmarkExecution, BenchmarkExecution.run_id == BenchmarkRun.id)
            .where(BenchmarkRun.is_deleted.is_(False))
            .group_by(BenchmarkRun.id)
            .order_by(BenchmarkRun.created_at.desc())
        )
        result = await db.execute(stmt)
        return list(result.all())

    async def get_run_details(self, db: AsyncSession, run_id: UUID) -> Optional[Tuple[BenchmarkRun, Dict[str, int]]]:
        stmt = (
            select(BenchmarkRun)
            .where(BenchmarkRun.id == run_id)
            .options(
                selectinload(BenchmarkRun.executions).joinedload(BenchmarkExecution.llm_model),
                selectinload(BenchmarkRun.executions).joinedload(BenchmarkExecution.test_case)
            )
        )
        result = await db.execute(stmt)
        run = result.scalar_one_or_none()

        if not run:
            return None

        status_counts = {
            ExecutionStatus.COMPLETED: 0,
            ExecutionStatus.FAILED: 0,
            ExecutionStatus.PENDING: 0,
            ExecutionStatus.PROCESSING: 0,
            ExecutionStatus.CANCELLED: 0
        }
        for e in run.executions:
            status_counts[e.status] += 1

        total = len(run.executions)
        completed = status_counts[ExecutionStatus.COMPLETED]
        failed = status_counts[ExecutionStatus.FAILED]
        processing = status_counts[ExecutionStatus.PROCESSING]
        pending = status_counts[ExecutionStatus.PENDING] + processing

        status_changed = False
        if processing > 0 and run.status != RunStatus.RUNNING:
            run.status = RunStatus.RUNNING
            status_changed = True
        elif total > 0 and (completed + failed) == total and run.status != RunStatus.COMPLETED:
            run.status = RunStatus.COMPLETED
            status_changed = True

        if status_changed:
            await db.commit()

        run.executions.sort(key=lambda e: (e.created_at, e.id))

        stats = {
            "total": total,
            "completed": completed,
            "failed": failed,
            "pending": pending
        }

        return run, stats

    async def delete_runs(self, db: AsyncSession, run_ids: List[UUID]) -> int:
        stmt = (
            update(BenchmarkRun)
            .where(BenchmarkRun.id.in_(run_ids))
            .values(is_deleted=True)
        )
        result = await db.execute(stmt)
        await db.commit()
        return result.rowcount  # type: ignore

    async def cancel_run(self, db: AsyncSession, run_id: UUID) -> bool:
        stmt = select(BenchmarkRun).where(BenchmarkRun.id == run_id)
        result = await db.execute(stmt)
        run = result.scalar_one_or_none()

        if not run or run.status in [RunStatus.COMPLETED, RunStatus.CANCELLED]:
            return False

        run.status = RunStatus.CANCELLED
        cancel_exec_stmt = (
            update(BenchmarkExecution)
            .where(
                BenchmarkExecution.run_id == run_id,
                BenchmarkExecution.status == ExecutionStatus.PENDING
            )
            .values(status=ExecutionStatus.CANCELLED)
        )
        await db.execute(cancel_exec_stmt)
        await db.commit()
        return True

    async def get_runs_summary(self, db: AsyncSession, run_ids: List[UUID]) -> list:
        stmt = (
            select(
                LLMModel.name.label("model_name"),
                TestSuite.name.label("suite_name"),
                func.avg(BenchmarkExecution.score).label("avg_score"),
                (func.avg(BenchmarkExecution.latency_ms) / 1000.0).label("avg_time"),
                func.avg(
                    func.coalesce(BenchmarkExecution.prompt_tokens, 0) +
                    func.coalesce(BenchmarkExecution.completion_tokens, 0)
                ).label("avg_tokens")
            )
            .join(LLMModel, BenchmarkExecution.llm_model_id == LLMModel.id)
            .join(TestCase, BenchmarkExecution.test_case_id == TestCase.id)
            .join(TestSuite, TestCase.suite_id == TestSuite.id)
            .where(BenchmarkExecution.run_id.in_(run_ids))
            .where(BenchmarkExecution.status.in_([ExecutionStatus.COMPLETED, ExecutionStatus.FAILED]))
            .group_by(LLMModel.name, TestSuite.name)
        )
        result = await db.execute(stmt)
        summary = result.all()

        return [
            {
                "model": row.model_name,
                "test_suite": row.suite_name,
                "avg_correctness": round((row.avg_score or 0) * 100, 2),
                "avg_time": round(row.avg_time or 0, 2),
                "avg_tokens": round(row.avg_tokens or 0, 0)
            }
            for row in summary
        ]

    async def export_to_excel(self, db: AsyncSession, run_ids: List[UUID]) -> bytes:
        if Workbook is None:  # pragma: no cover
            raise RuntimeError("openpyxl is not installed.")

        summary = await self.get_runs_summary(db, run_ids)
        excel_bytes = await asyncio.to_thread(self._generate_excel_sync, summary)
        return excel_bytes

    def _generate_excel_sync(self, summary_data: list) -> bytes:
        rows = sorted([
            [
                item["model"],
                item["test_suite"],
                item["avg_correctness"],
                item["avg_time"],
                item["avg_tokens"]
            ] for item in summary_data
        ], key=lambda x: (x[0], x[1]))

        wb = Workbook()
        ws = wb.active
        ws.title = "Benchmark Results"  # type: ignore

        headers = ["Model", "Zbiór testowy", "Średnia poprawność", "Średni czas procesowania", "Średnie zużycie tokenów"]
        ws.append(headers)  # type: ignore

        for row in rows:
            ws.append(row)  # type: ignore

        header_font = Font(bold=True)
        for cell in ws[1]:  # type: ignore
            cell.font = header_font

        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):  # type: ignore
            for cell in row:
                cell.border = thin_border

        for idx, column_cells in enumerate(ws.columns, 1):  # type: ignore
            length = max(len(str(cell.value)) for cell in column_cells)
            ws.column_dimensions[get_column_letter(idx)].width = min(length + 2, 50)  # type: ignore

        excel_bytes = io.BytesIO()
        wb.save(excel_bytes)
        excel_bytes.seek(0)
        return excel_bytes.read()

    async def repeat_execution(self, db: AsyncSession, execution_id: UUID) -> Optional[BenchmarkExecution]:
        stmt = (
            update(BenchmarkExecution)
            .where(BenchmarkExecution.id == execution_id)
            .values(
                status=ExecutionStatus.PENDING,
                response_text=None,
                score=None,
                error_message=None,
                prompt_tokens=None,
                completion_tokens=None,
                latency_ms=None
            )
            .returning(BenchmarkExecution)
        )
        result = await db.execute(stmt)
        execution = result.scalar_one_or_none()

        if execution:
            await db.commit()

        return execution

    async def repeat_failed_executions(self, db: AsyncSession, run_id: UUID) -> Optional[Tuple[BenchmarkRun, Dict[str, int]]]:
        reset_stmt = (
            update(BenchmarkExecution)
            .where(
                BenchmarkExecution.run_id == run_id,
                BenchmarkExecution.status == ExecutionStatus.FAILED
            )
            .values(
                status=ExecutionStatus.PENDING,
                response_text=None,
                score=None,
                error_message=None,
                prompt_tokens=None,
                completion_tokens=None,
                latency_ms=None
            )
        )
        result = await db.execute(reset_stmt)

        if result.rowcount > 0:  # type: ignore
            update_run_stmt = update(BenchmarkRun).where(BenchmarkRun.id == run_id).values(status=RunStatus.PENDING)
            await db.execute(update_run_stmt)
            await db.commit()

        return await self.get_run_details(db, run_id)


benchmark_service = BenchmarkService()
